from __future__ import annotations

import argparse
import math
import re
import statistics
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

"""
02_pv.neutralisation.graphs.py

Purpose:
    Fit 4-parameter logistic curves to pseudoneutralisation data and generate diagnostic plots.

Input:
    Excel workbook containing pseudoneutralisation data.

Output:
    1. Excel workbook containing summary of pVNT50, pVNT90, and Hill slope values for each sample/pseudovirus combination.
	2. PNG files containing pseudoneutralisation curves with 4PL fits for each sample/pseudovirus combination.

Thesis:
    Chapter: Recombinant virology toolkit for ERVEV

Author:
    Igor Starinskij

AI assistance:
    Generative artificial intelligence tools, primarily GitHub Copilot 0.63.0 (GPT-5.6 Luna), were used during the development of software used in this thesis.
	AI was used to assist with Python code generation, debugging, code refinement and explanation of programming concepts.
	Generated code was reviewed, modified, tested and validated by the author.
	The analytical methodology, statistical decisions, interpretation of results and final responsibility for the analyses remain with the author.
"""

NEG_CONTROL_LABEL = "neg"
POS_CONTROL_LABEL = "pos"
DEFAULT_OUTPUT_NAME = "pv_neutralisation_results.xlsx"


@dataclass
class NeutralisationSeries:
	pseudovirus: str
	sample: str
	columns: List[int]


@dataclass
class NeutralisationResult:
	pseudovirus: str
	sample: str
	pvnt50: Optional[float]
	pvnt90: Optional[float]
	hill_slope: Optional[float]
	status: str
	points_used: int


def normalize_text(value: object) -> str:
	if value is None:
		return ""
	return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def safe_filename_component(value: object) -> str:
	text = normalize_text(value)
	text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
	text = text.strip("._-")
	return text or "unknown"


def parse_dilution_denominator(value: object) -> Optional[float]:
	if value is None:
		return None
	if isinstance(value, (int, float)):
		if float(value) <= 0:
			return None
		return float(value)
	text = normalize_text(value)
	if not text:
		return None
	match = re.search(r"1\s*[:/]\s*(\d+(?:\.\d+)?)", text)
	if match:
		try:
			denominator = float(match.group(1))
		except ValueError:
			return None
		return denominator if denominator > 0 else None
	try:
		denominator = float(text)
	except ValueError:
		return None
	return denominator if denominator > 0 else None


def parse_header(header: object) -> Optional[Tuple[str, str]]:
	text = normalize_text(header)
	if not text:
		return None
	parts = text.split(None, 1)
	pseudovirus = parts[0].strip()
	sample = parts[1].strip() if len(parts) > 1 else ""
	return pseudovirus, sample


def control_kind(sample: object) -> Optional[str]:
	text = normalize_text(sample).lower()
	if not text:
		return None
	first_token = text.split(None, 1)[0].strip("[]")
	if first_token == POS_CONTROL_LABEL or first_token.startswith(f"{POS_CONTROL_LABEL} "):
		return POS_CONTROL_LABEL
	if first_token == NEG_CONTROL_LABEL or first_token.startswith(f"{NEG_CONTROL_LABEL} "):
		return NEG_CONTROL_LABEL
	if re.search(r"\bpos\b", text):
		return POS_CONTROL_LABEL
	if re.search(r"\bneg\b", text):
		return NEG_CONTROL_LABEL
	return None


def control_header_parts(header: object) -> Optional[Tuple[Optional[str], str]]:
	text = normalize_text(header)
	if not text:
		return None
	kind = control_kind(text)
	if kind is None:
		return None
	parts = text.split()
	if kind == NEG_CONTROL_LABEL:
		if len(parts) == 1:
			return None, NEG_CONTROL_LABEL
		if parts[-1].strip("[]").lower() == NEG_CONTROL_LABEL:
			return " ".join(parts[:-1]), NEG_CONTROL_LABEL
		return None, NEG_CONTROL_LABEL
	if parts[-1].strip("[]").lower() == POS_CONTROL_LABEL and len(parts) >= 2:
		return " ".join(parts[:-1]), POS_CONTROL_LABEL
	return None


def group_series_columns(ws) -> Dict[Tuple[str, str], NeutralisationSeries]:
	series: Dict[Tuple[str, str], NeutralisationSeries] = {}
	for col in range(2, ws.max_column + 1):
		if control_header_parts(ws.cell(1, col).value) is not None:
			continue
		header = parse_header(ws.cell(1, col).value)
		if header is None:
			continue
		pseudovirus, sample = header
		key = (pseudovirus.upper(), sample.lower())
		group = series.get(key)
		if group is None:
			group = NeutralisationSeries(pseudovirus=pseudovirus.upper(), sample=sample, columns=[])
			series[key] = group
		group.columns.append(col)
	return series


def mean_sd(values: Sequence[float]) -> Tuple[Optional[float], Optional[float]]:
	clean = [float(value) for value in values if value is not None]
	if not clean:
		return None, None
	if len(clean) == 1:
		return clean[0], 0.0
	return statistics.mean(clean), statistics.stdev(clean)


def logistic_4pl_inhibition(x_values, bottom: float, top: float, log_ic50: float, hill_slope: float):
	import numpy as np

	x_values = np.asarray(x_values, dtype=float)
	exponent = np.clip((x_values - log_ic50) * hill_slope, -20.0, 20.0)
	return bottom + (top - bottom) / (1.0 + np.power(10.0, exponent))


def fit_4pl_inhibition(x_values: Sequence[float], y_values: Sequence[float]):
	try:
		import numpy as np
		from scipy.optimize import curve_fit
	except Exception:
		curve_fit = None
		import numpy as np

	x = np.asarray(list(x_values), dtype=float)
	y = np.asarray(list(y_values), dtype=float)
	if x.size < 4 or y.size < 4:
		return None
	if not np.all(np.isfinite(x)) or not np.all(np.isfinite(y)):
		return None
	if float(np.ptp(y)) < 1e-6:
		return None

	x_min = float(np.min(x))
	x_max = float(np.max(x))
	y_min = float(np.min(y))
	y_max = float(np.max(y))

	bottom_seed = max(0.0, min(100.0, y_min))
	top_seed = max(bottom_seed + 1.0, min(100.0, y_max))
	log_ic50_seed = (x_min + x_max) / 2.0
	hill_seed = 1.0
	p0 = [bottom_seed, top_seed, log_ic50_seed, hill_seed]
	bounds = ([0.0, 0.0, x_min - 3.0, 0.01], [100.0, 100.0, x_max + 3.0, 10.0])

	if curve_fit is not None:
		try:
			with warnings.catch_warnings():
				warnings.simplefilter("ignore")
				params, _ = curve_fit(
					logistic_4pl_inhibition,
					x,
					y,
					p0=p0,
					bounds=bounds,
					maxfev=20000,
				)
			return tuple(float(value) for value in params)
		except Exception:
			pass

	def sse(params: Sequence[float]) -> float:
		predicted = logistic_4pl_inhibition(x, *params)
		if not np.all(np.isfinite(predicted)):
			return float("inf")
		return float(np.sum((predicted - y) ** 2))

	best = list(p0)
	best_score = sse(best)
	bottom_grid = np.linspace(max(0.0, y_min - 20.0), min(95.0, y_min + 20.0), 5)
	top_grid = np.linspace(max(bottom_grid[-1] + 1.0, y_max - 20.0), min(100.0, y_max + 20.0), 5)
	center_grid = np.linspace(x_min, x_max, 7)
	hill_grid = np.array([0.25, 0.5, 1.0, 2.0, 4.0])
	for bottom in bottom_grid:
		for top in top_grid:
			if top <= bottom:
				continue
			for center in center_grid:
				for hill in hill_grid:
					candidate = [float(bottom), float(top), float(center), float(hill)]
					score = sse(candidate)
					if score < best_score:
						best_score = score
						best = candidate

	steps = [10.0, 10.0, max(0.15, (x_max - x_min) / 6.0), 0.5]
	for _ in range(60):
		improved = False
		for idx in range(4):
			for direction in (-1.0, 1.0):
				candidate = best.copy()
				candidate[idx] = max(bounds[0][idx], min(bounds[1][idx], candidate[idx] + direction * steps[idx]))
				score = sse(candidate)
				if score < best_score:
					best_score = score
					best = candidate
					improved = True
		if not improved:
			steps = [step * 0.5 for step in steps]
			if max(steps) < 1e-3:
				break

	return tuple(float(value) for value in best)


def inverse_4pl_inhibition(params: Sequence[float], threshold: float) -> Optional[float]:
	if params is None or len(params) != 4:
		return None
	bottom, top, log_ic50, hill_slope = (float(value) for value in params)
	if hill_slope <= 0 or top <= bottom:
		return None
	if threshold <= bottom or threshold >= top:
		return None
	ratio = (top - bottom) / (threshold - bottom) - 1.0
	if ratio <= 0:
		return None
	x_value = log_ic50 + (math.log10(ratio) / hill_slope)
	return 10 ** x_value


def choose_input_file() -> Optional[Path]:
	from tkinter import Tk, filedialog

	root = Tk()
	root.withdraw()
	selected = filedialog.askopenfilename(
		title="Select pseudoneutralisation workbook",
		filetypes=[("Excel workbooks", "*.xlsx *.xlsm *.xltx *.xltm"), ("All files", "*.*")],
	)
	root.destroy()
	if not selected:
		return None
	return Path(selected)


def choose_output_directory() -> Optional[Path]:
	from tkinter import Tk, filedialog

	root = Tk()
	root.withdraw()
	selected = filedialog.askdirectory(title="Choose output folder")
	root.destroy()
	if not selected:
		return None
	output_dir = Path(selected)
	output_dir.mkdir(parents=True, exist_ok=True)
	return output_dir


def choose_data_sheet(wb, sheet_name: Optional[str] = None):
	if sheet_name:
		if sheet_name not in wb.sheetnames:
			raise ValueError(f"Workbook does not contain sheet {sheet_name!r}")
		return wb[sheet_name]
	best_sheet = None
	best_score = -1
	for ws in wb.worksheets:
		score = 0
		first_row_values = [normalize_text(ws.cell(1, col).value).lower() for col in range(1, ws.max_column + 1)]
		if any(value in {"dilution", "dil", "serum dilution"} for value in first_row_values[:3]):
			score += 4
		if any(parse_header(ws.cell(1, col).value) is not None for col in range(2, min(ws.max_column, 20) + 1)):
			score += 2
		if any(control_kind(ws.cell(1, col).value) is not None for col in range(2, ws.max_column + 1)):
			score += 2
		if ws.max_row >= 4 and ws.max_column >= 4:
			score += 1
		if score > best_score:
			best_score = score
			best_sheet = ws
	if best_sheet is not None and best_score > 0:
		return best_sheet
	for ws in wb.worksheets:
		if ws.max_row >= 2 and ws.max_column >= 2:
			return ws
	raise ValueError("Workbook does not contain a usable data sheet")


def collect_control_values(ws) -> Dict[Tuple[str, str], List[float]]:
	values_by_control: Dict[Tuple[str, str], List[float]] = {}
	for col in range(2, ws.max_column + 1):
		header_parts = control_header_parts(ws.cell(1, col).value)
		if header_parts is None:
			continue
		pseudovirus, kind = header_parts
		key = ((pseudovirus or "").upper(), kind)
		if key == ("", NEG_CONTROL_LABEL):
			key = ("GLOBAL", NEG_CONTROL_LABEL)
		for row in range(2, ws.max_row + 1):
			cell_value = ws.cell(row, col).value
			if cell_value in (None, ""):
				continue
			try:
				values_by_control.setdefault(key, []).append(float(cell_value))
			except (TypeError, ValueError):
				continue
	return values_by_control


def collect_series_values(ws, columns: Sequence[int]) -> Dict[float, List[float]]:
	values_by_dilution: Dict[float, List[float]] = {}
	for row in range(2, ws.max_row + 1):
		dilution = parse_dilution_denominator(ws.cell(row, 1).value)
		if dilution is None:
			continue
		row_values: List[float] = []
		for col in columns:
			cell_value = ws.cell(row, col).value
			if cell_value in (None, ""):
				continue
			try:
				row_values.append(float(cell_value))
			except (TypeError, ValueError):
				continue
		if row_values:
			values_by_dilution.setdefault(dilution, []).extend(row_values)
	return values_by_dilution


def collect_row_means(ws, columns: Sequence[int]) -> Dict[float, List[float]]:
	values_by_dilution: Dict[float, List[float]] = {}
	for row in range(2, ws.max_row + 1):
		dilution = parse_dilution_denominator(ws.cell(row, 1).value)
		if dilution is None:
			continue
		row_values: List[float] = []
		for col in columns:
			cell_value = ws.cell(row, col).value
			if cell_value in (None, ""):
				continue
			try:
				row_values.append(float(cell_value))
			except (TypeError, ValueError):
				continue
		if row_values:
			values_by_dilution.setdefault(dilution, []).append(statistics.mean(row_values))
	return values_by_dilution


def control_mean(values_by_control: Dict[Tuple[str, str], List[float]], pseudovirus: str, kind: str) -> Optional[float]:
	values = values_by_control.get((pseudovirus.upper(), kind))
	if not values and kind == NEG_CONTROL_LABEL:
		values = values_by_control.get(("GLOBAL", NEG_CONTROL_LABEL))
	if not values:
		return None
	return statistics.mean(values)


def build_results(ws) -> Tuple[List[NeutralisationResult], Dict[Tuple[str, str], Dict[str, object]]]:
	series_groups = group_series_columns(ws)
	if not series_groups:
		raise ValueError("No sample/control headers found in the first row")

	control_values = collect_control_values(ws)

	results: List[NeutralisationResult] = []
	plot_payload: Dict[Tuple[str, str], Dict[str, object]] = {}

	for series in series_groups.values():
		label = control_kind(series.sample)
		if label in {POS_CONTROL_LABEL, NEG_CONTROL_LABEL}:
			continue

		pos_mean = control_mean(control_values, series.pseudovirus, POS_CONTROL_LABEL)
		neg_mean = control_mean(control_values, series.pseudovirus, NEG_CONTROL_LABEL)
		if pos_mean is None or neg_mean is None:
			# Keep the series in the output so missing controls are visible instead of silently dropped.
			results.append(
				NeutralisationResult(
					pseudovirus=series.pseudovirus,
					sample=series.sample,
					pvnt50=None,
					pvnt90=None,
					hill_slope=None,
					status="missing controls",
					points_used=0,
				)
			)
			continue

		sample_by_dilution = collect_series_values(ws, series.columns)

		ordered_dilutions = sorted(sample_by_dilution.keys(), key=lambda dil: 1.0 / dil, reverse=True)
		if not ordered_dilutions:
			results.append(
				NeutralisationResult(
					pseudovirus=series.pseudovirus,
					sample=series.sample,
					pvnt50=None,
					pvnt90=None,
					hill_slope=None,
					status="no usable points",
					points_used=0,
				)
			)
			continue

		neutralisation_means: List[float] = []
		neutralisation_sds: List[float] = []
		sample_means: List[float] = []
		sample_sds: List[float] = []
		pos_means: List[float] = []
		pos_sds: List[float] = []
		neg_means: List[float] = []
		neg_sds: List[float] = []

		for dilution in ordered_dilutions:
			sample_values = sample_by_dilution[dilution]
			sample_mean, sample_sd = mean_sd(sample_values)
			if sample_mean is None:
				continue
			denominator = pos_mean - neg_mean
			if denominator == 0:
				continue
			neutralisation_values = [
				100.0 * (pos_mean - value) / denominator
				for value in sample_values
				if value is not None
			]
			if not neutralisation_values:
				continue
			neutralisation_mean, neutralisation_sd = mean_sd(neutralisation_values)
			if neutralisation_mean is None:
				continue
			neutralisation_means.append(max(0.0, min(100.0, neutralisation_mean)))
			neutralisation_sds.append(neutralisation_sd or 0.0)
			sample_means.append(sample_mean)
			sample_sds.append(sample_sd or 0.0)
			pos_means.append(pos_mean)
			pos_sds.append(0.0)
			neg_means.append(neg_mean)
			neg_sds.append(0.0)

		if len(neutralisation_means) < 4:
			results.append(
				NeutralisationResult(
					pseudovirus=series.pseudovirus,
					sample=series.sample,
					pvnt50=None,
					pvnt90=None,
					hill_slope=None,
					status="insufficient points",
					points_used=len(neutralisation_means),
				)
			)
			plot_payload[(series.pseudovirus, series.sample)] = {
				"ordered_dilutions": ordered_dilutions,
				"neutralisation_means": neutralisation_means,
				"neutralisation_sds": neutralisation_sds,
				"sample_means": sample_means,
				"sample_sds": sample_sds,
				"pos_means": pos_means,
				"pos_sds": pos_sds,
				"neg_means": neg_means,
				"neg_sds": neg_sds,
				"pos_mean": pos_mean,
				"neg_mean": neg_mean,
				"pVNT50": None,
				"pVNT90": None,
				"hill_slope": None,
			}
			continue

		fit_params = fit_4pl_inhibition(
			[math.log10(1.0 / dil) for dil in ordered_dilutions[: len(neutralisation_means)]],
			neutralisation_means,
		)
		pVNT90 = inverse_4pl_inhibition(fit_params, 90.0) if fit_params is not None else None
		pVNT50 = inverse_4pl_inhibition(fit_params, 50.0) if fit_params is not None else None
		hill_slope = float(fit_params[3]) if fit_params is not None else None

		results.append(
			NeutralisationResult(
				pseudovirus=series.pseudovirus,
				sample=series.sample,
				pvnt50=pVNT50,
				pvnt90=pVNT90,
				hill_slope=hill_slope,
				status="ok",
				points_used=len(neutralisation_means),
			)
		)
		plot_payload[(series.pseudovirus, series.sample)] = {
			"ordered_dilutions": ordered_dilutions[: len(neutralisation_means)],
			"neutralisation_means": neutralisation_means,
			"neutralisation_sds": neutralisation_sds,
			"sample_means": sample_means,
			"sample_sds": sample_sds,
			"pos_means": pos_means,
			"pos_sds": pos_sds,
			"neg_means": neg_means,
			"neg_sds": neg_sds,
			"pos_mean": pos_mean,
			"neg_mean": neg_mean,
			"pVNT50": pVNT50,
			"pVNT90": pVNT90,
			"hill_slope": hill_slope,
		}

	return results, plot_payload


def write_summary_workbook(output_path: Path, results: Sequence[NeutralisationResult]) -> None:
	wb = Workbook()
	ws = wb.active
	ws.title = "results"
	headers = ["pseudovirus", "sample", "pVNT50", "pVNT90", "hill_slope", "points_used", "status"]
	for col, header in enumerate(headers, start=1):
		cell = ws.cell(1, col, header)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal="center")

	for row, result in enumerate(sorted(results, key=lambda item: (item.pseudovirus, item.sample)), start=2):
		ws.cell(row, 1, result.pseudovirus)
		ws.cell(row, 2, result.sample)
		ws.cell(row, 3, round(result.pvnt50) if result.pvnt50 is not None else None)
		ws.cell(row, 4, round(result.pvnt90) if result.pvnt90 is not None else None)
		ws.cell(row, 5, round(result.hill_slope, 6) if result.hill_slope is not None else None)
		ws.cell(row, 6, result.points_used)
		ws.cell(row, 7, result.status)

	for column, width in {1: 16, 2: 34, 3: 14, 4: 14, 5: 12, 6: 12, 7: 20}.items():
		ws.column_dimensions[chr(64 + column)].width = width

	wb.save(output_path)


def plot_results(output_dir: Path, results: Sequence[NeutralisationResult], payload: Dict[Tuple[str, str], Dict[str, object]]) -> None:
	try:
		import matplotlib
		matplotlib.use("Agg")
		import matplotlib.pyplot as plt
		import numpy as np
	except Exception as exc:  # pragma: no cover - dependency/runtime guard
		raise RuntimeError("matplotlib is required to generate PV graphs") from exc

	for result in results:
		data = payload.get((result.pseudovirus, result.sample))
		if not data:
			continue
		ordered_dilutions = list(data["ordered_dilutions"])
		neutralisation_means = list(data["neutralisation_means"])
		neutralisation_sds = list(data["neutralisation_sds"])
		sample_means = list(data["sample_means"])
		sample_sds = list(data["sample_sds"])
		pos_mean = data.get("pos_mean")
		neg_mean = data.get("neg_mean")
		fit_params = None
		if result.pvnt50 is not None or result.pvnt90 is not None:
			fit_params = fit_4pl_inhibition(
				[math.log10(1.0 / dil) for dil in ordered_dilutions],
				neutralisation_means,
			)

		x_positions = list(range(len(ordered_dilutions) + 2))
		x_labels = [f"1:{int(round(1.0 / dil))}" for dil in ordered_dilutions]

		fig, ax = plt.subplots(figsize=(8.5, 5.5), dpi=200)
		fig.subplots_adjust(right=0.88)
		lum_line = None

		neutralisation_container = ax.errorbar(
			list(range(1, 1 + len(ordered_dilutions))),
			neutralisation_means,
			yerr=neutralisation_sds,
			fmt="o",
			markersize=4.5,
			linewidth=0.0,
			color="#0b6efd",
			ecolor="#0b6efd",
			elinewidth=1.0,
			capsize=3,
			label="% neutralisation",
			zorder=4,
		)
		if fit_params is not None:
			fit_x = np.linspace(math.log10(1.0 / ordered_dilutions[-1]), math.log10(1.0 / ordered_dilutions[0]), 200)
			fit_x_plot = np.interp(fit_x, [math.log10(1.0 / dil) for dil in ordered_dilutions[::-1]], list(range(len(ordered_dilutions), 0, -1)))
			fit_y = logistic_4pl_inhibition(fit_x, *fit_params)
			ax.plot(fit_x_plot, fit_y, linewidth=2.2, color="#0b6efd", label="_nolegend_", zorder=3)

		ax.set_ylim(0, 100)
		ax.set_yticks([0, 25, 50, 75, 100])
		ax.set_ylabel("% neutralisation")
		ax.set_xlabel("Dilution")
		ax.set_xticks(x_positions)
		ax.set_xticklabels(["No serum control", *x_labels, "Negative control"], rotation=45, ha="right")
		ax.grid(True, axis="y", alpha=0.25)
		ax.set_title(
			f"{result.sample} {result.pseudovirus} pseudoneutralisation curve with 4PL fit",
			loc="left",
			fontsize=13,
			pad=22,
		)

		ax2 = ax.twinx()
		ax2.set_yscale("log")
		max_lum_value = max(float(pos_mean or 1.0), max(sample_means, default=1.0), float(neg_mean or 1.0))
		lum_upper = 10 ** 7 if max_lum_value > 10 ** 6 else 10 ** 6
		ax2.set_ylim(1, lum_upper)
		ax2.set_ylabel("Luminescence (cps)", labelpad=2)
		ax2.tick_params(axis="y", colors="#d55e00")
		ax2.spines["right"].set_color("#d55e00")

		if pos_mean is not None:
			ax2.errorbar(
				[x_positions[0]],
				[pos_mean],
				yerr=[[0.0], [0.0]],
				fmt="s",
				markersize=4.5,
				linewidth=0.0,
				color="#d55e00",
				ecolor="#d55e00",
				elinewidth=1.0,
				capsize=3,
				label="_nolegend_",
				zorder=3,
			)
		if sample_means:
			lum_line = ax2.errorbar(
				x_positions[1:-1],
				sample_means,
				yerr=[sample_sds, sample_sds],
				fmt="s-",
				markersize=4.5,
				linewidth=2.0,
				color="#d55e00",
				ecolor="#d55e00",
				elinewidth=1.0,
				capsize=3,
				label="Luminescence",
				zorder=3,
			)
		if neg_mean is not None:
			ax2.errorbar(
				[x_positions[-1]],
				[neg_mean],
				yerr=[[0.0], [0.0]],
				fmt="s",
				markersize=4.5,
				linewidth=0.0,
				color="#d55e00",
				ecolor="#d55e00",
				elinewidth=1.0,
				capsize=3,
				label="_nolegend_",
				zorder=3,
			)

		fit_text = [
			f"pVNT50: {int(round(result.pvnt50))}" if result.pvnt50 is not None else "pVNT50: NA",
			f"pVNT90: {int(round(result.pvnt90))}" if result.pvnt90 is not None else "pVNT90: NA",
			f"Hill slope: {result.hill_slope:.3g}" if result.hill_slope is not None else "Hill slope: NA",
		]
		info_lines = [
			f"Sample: {result.sample}",
			f"PV: {result.pseudovirus}",
			f"Points used: {result.points_used}",
			*fit_text,
			f"Status: {result.status}",
		]
		info_width = 34
		info_text = "\n".join(f"{line:^{info_width}}" for line in info_lines)
		fig.text(
			0.99,
			0.87,
			info_text,
			ha="left",
			va="top",
			fontsize=9,
			bbox={"boxstyle": "round,pad=0.55", "facecolor": "white", "edgecolor": "#999999", "alpha": 0.92},
		)
		handles = [neutralisation_container]
		labels = ["% neutralisation"]
		if lum_line is not None:
			handles.append(lum_line)
			labels.append("Luminescence")
		fig.legend(handles, labels, loc="upper left", bbox_to_anchor=(0.99, 0.63), frameon=False)
		ax2.grid(False)

		series_dir = output_dir / safe_filename_component(result.pseudovirus)
		series_dir.mkdir(parents=True, exist_ok=True)
		filename = f"{safe_filename_component(result.sample)}_{safe_filename_component(result.pseudovirus)}.png"
		fig.savefig(series_dir / filename, bbox_inches="tight")
		plt.close(fig)


def process_workbook(input_path: Path, output_dir: Path, sheet_name: Optional[str] = None) -> Path:
	wb = load_workbook(input_path, data_only=True)
	ws = choose_data_sheet(wb, sheet_name=sheet_name)
	results, payload = build_results(ws)

	output_dir.mkdir(parents=True, exist_ok=True)
	summary_path = output_dir / DEFAULT_OUTPUT_NAME
	write_summary_workbook(summary_path, results)
	plot_results(output_dir, results, payload)
	return summary_path


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description=(
			"Create pseudoneutralisation graphs from a workbook where the first row contains"
			" '<pseudovirus> <sample>' headers and the first column contains serum dilutions."
		)
	)
	parser.add_argument("input", nargs="?", type=Path, help="Input Excel workbook")
	parser.add_argument("--output-dir", type=Path, help="Folder for graphs and summary workbook")
	parser.add_argument("--sheet", help="Worksheet name to read, if the workbook has multiple sheets")
	return parser.parse_args()


def main() -> None:
	args = parse_args()
	input_path = args.input or choose_input_file()
	if input_path is None:
		return
	if not input_path.exists():
		raise FileNotFoundError(f"Input workbook does not exist: {input_path}")

	output_dir = args.output_dir
	if output_dir is None:
		output_dir = choose_output_directory()
		if output_dir is None:
			return

	summary_path = process_workbook(input_path, output_dir, sheet_name=args.sheet)

	try:
		from tkinter import Tk, messagebox

		root = Tk()
		root.withdraw()
		messagebox.showinfo("PV neutralisation complete", f"Saved summary workbook to:\n{summary_path}")
		root.destroy()
	except Exception:
		print(f"Saved summary workbook to: {summary_path}")


if __name__ == "__main__":
	main()