"""Run univariable and multivariable logistic regression from an Excel sheet."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

"""
09_multivariate.stepwise.regression.py

Purpose:
    Run univariable and multivariable logistic regression from an Excel sheet.

Input:
    Excel workbook containing a sheet named "all" with a binary response variable and one or more explanatory variables.


Output:
    Excel workbook containing a summary of univariable and multivariable logistic regression results for each explanatory variable, including odds ratios, confidence intervals, p-values, omnibus p-values, and model quality metrics.

Thesis:
    Chapter: ERVEV biosurveillance survey

Author:
    Igor Starinskij

AI assistance:
    Generative artificial intelligence tools, primarily GitHub Copilot 0.63.0 (GPT-5.6 Luna), were used during the development of software used in this thesis.
	AI was used to assist with Python code generation, debugging, code refinement and explanation of programming concepts.
	Generated code was reviewed, modified, tested and validated by the author.
	The analytical methodology, statistical decisions, interpretation of results and final responsibility for the analyses remain with the author.
"""

def choose_file() -> Path | None:
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    selected = filedialog.askopenfilename(
        title="Select Excel workbook",
        filetypes=[("Excel workbooks", "*.xlsx"), ("All files", "*.*")],
    )
    root.destroy()
    return Path(selected) if selected else None


def choose_columns(columns: list[str], title: str, multiple: bool) -> list[str]:
    import tkinter as tk
    from tkinter import messagebox

    result: list[str] = []
    root = tk.Tk()
    root.title(title)
    root.resizable(False, False)
    tk.Label(root, text="Select one or more columns, then click OK.").pack(
        padx=12, pady=(12, 4)
    )
    listbox = tk.Listbox(
        root,
        selectmode=tk.EXTENDED if multiple else tk.SINGLE,
        width=48,
        height=min(25, max(5, len(columns))),
        exportselection=False,
    )
    for column in columns:
        listbox.insert(tk.END, column)
    listbox.pack(padx=12, pady=4)

    def accept() -> None:
        selected = listbox.curselection()
        if not selected:
            messagebox.showerror("Selection required", "Select at least one column.", parent=root)
            return
        result.extend(columns[index] for index in selected)
        root.destroy()

    tk.Button(root, text="OK", command=accept, width=12).pack(pady=(4, 12))
    root.protocol("WM_DELETE_WINDOW", root.destroy)
    root.mainloop()
    return result


def choose_references(data: pd.DataFrame, predictors: list[str]) -> dict[str, Any]:
    import tkinter as tk
    from tkinter import messagebox

    references: dict[str, Any] = {}
    categorical = [predictor for predictor in predictors if not _is_numeric(data[predictor].dropna())]
    for predictor in categorical:
        levels = list(pd.unique(data[predictor].dropna().astype(str)))
        if len(levels) < 2:
            continue
        selected = [levels[0]]
        root = tk.Tk()
        root.title(f"Reference level: {predictor}")
        root.resizable(False, False)
        tk.Label(root, text=f"Select the reference level for '{predictor}'.").pack(
            padx=12, pady=(12, 4)
        )
        listbox = tk.Listbox(root, selectmode=tk.SINGLE, width=48, height=min(20, len(levels)))
        for level in levels:
            listbox.insert(tk.END, level)
        listbox.selection_set(0)
        listbox.pack(padx=12, pady=4)

        def accept() -> None:
            selection = listbox.curselection()
            if selection:
                selected[0] = levels[selection[0]]
            root.destroy()

        tk.Button(root, text="OK", command=accept, width=12).pack(pady=(4, 12))
        root.protocol("WM_DELETE_WINDOW", root.destroy)
        root.mainloop()
        references[predictor] = selected[0]
    return references


def load_binary_data(path: Path, response: str, predictors: list[str]) -> pd.DataFrame:
    data = pd.read_excel(path, sheet_name="all")
    data.columns = [str(column).strip() for column in data.columns]
    missing = [column for column in [response, *predictors] if column not in data.columns]
    if missing:
        raise ValueError("Column(s) not found: " + ", ".join(missing))

    selected = data[[response, *predictors]].copy()
    numeric_response = pd.to_numeric(selected[response], errors="coerce")
    invalid = numeric_response.notna() & ~numeric_response.isin([0, 1])
    if invalid.any() or numeric_response.dropna().nunique() != 2:
        raise ValueError(f"'{response}' must contain both 0 and 1, with no other numeric values.")
    selected[response] = numeric_response
    return selected


def _is_numeric(series: pd.Series) -> bool:
    converted = pd.to_numeric(series, errors="coerce")
    return converted.notna().sum() == series.notna().sum()


def make_design(
    data: pd.DataFrame,
    response: str,
    predictors: list[str],
    references: dict[str, Any] | None = None,
) -> tuple[pd.DataFrame, pd.Series, dict[str, dict[str, Any]]]:
    references = references or {}
    complete = data[[response, *predictors]].dropna().copy()
    if complete.empty or complete[response].nunique() < 2:
        raise ValueError("The selected data do not contain complete rows with both response values.")

    design = pd.DataFrame(index=complete.index)
    details: dict[str, dict[str, Any]] = {}
    for predictor in predictors:
        values = complete[predictor]
        if _is_numeric(values):
            design[predictor] = pd.to_numeric(values)
            details[predictor] = {"numeric": True, "levels": list(pd.unique(values))}
            continue
        categorical = values.astype(str)
        levels = list(pd.unique(categorical))
        reference = str(references.get(predictor, levels[0]))
        if reference not in levels:
            reference = levels[0]
        ordered_levels = [reference, *[level for level in levels if level != reference]]
        details[predictor] = {"numeric": False, "levels": ordered_levels, "reference": reference}
        for level in ordered_levels[1:]:
            design[f"{predictor}={level}"] = (categorical == level).astype(float)

    return design.astype(float), complete[response].astype(float), details


def fit_logistic(
    design: pd.DataFrame, response: pd.Series
) -> tuple[dict[str, tuple[float, float, float, float]], float | None, set[str], float | None, int, dict[str, Any] | None]:
    try:
        import statsmodels.api as sm
    except ImportError as error:
        raise RuntimeError(
            "This script requires statsmodels. Install it with: python -m pip install statsmodels"
        ) from error

    if design.shape[1] == 0:
        return {}, None, set(), None, 0, None

    active = design.copy()
    excluded: set[str] = set()
    for name in list(active.columns):
        if not set(active[name].dropna().unique()).issubset({0.0, 1.0}):
            continue
        exposed = response[active[name] == 1]
        unexposed = response[active[name] == 0]
        if len(exposed) and len(unexposed) and (
            exposed.nunique() == 1 or unexposed.nunique() == 1
        ):
            excluded.add(name)
    active = active.drop(columns=excluded)
    while active.shape[1]:
        model = sm.GLM(response, sm.add_constant(active, has_constant="add"), family=sm.families.Binomial())
        try:
            fitted = model.fit()
        except Exception:
            name = active.columns[-1]
            excluded.add(name)
            active = active.drop(columns=[name])
            continue
        if np.isfinite(fitted.params[active.columns]).all() and np.abs(fitted.params[active.columns]).max() < 25:
            break
        name = fitted.params[active.columns].abs().idxmax()
        excluded.add(name)
        active = active.drop(columns=[name])
    if not active.shape[1]:
        return {}, None, excluded, None, 0, None

    model = sm.GLM(response, sm.add_constant(active, has_constant="add"), family=sm.families.Binomial())
    fitted = model.fit()
    confidence = fitted.conf_int()
    results: dict[str, tuple[float, float, float, float]] = {}
    for name in active.columns:
        coefficient = fitted.params[name]
        results[name] = (
            float(np.exp(coefficient)),
            float(np.exp(confidence.loc[name, 0])),
            float(np.exp(confidence.loc[name, 1])),
            float(fitted.pvalues[name]),
        )
    null_model = sm.GLM(response, np.ones((len(response), 1)), family=sm.families.Binomial()).fit()
    from scipy.stats import chi2

    omnibus = float(chi2.sf(2 * (fitted.llf - null_model.llf), len(active.columns)))
    probabilities = np.asarray(fitted.predict(), dtype=float)
    order = np.argsort(probabilities)
    ranked_response = response.to_numpy()[order]
    positives = ranked_response.sum()
    negatives = len(ranked_response) - positives
    auc = float(
        (np.flatnonzero(ranked_response == 1) + 1).sum() - positives * (positives + 1) / 2
    ) / float(positives * negatives) if positives and negatives else None
    auc_ci_low = auc_ci_high = None
    if auc is not None:
        q1 = auc / (2 - auc)
        q2 = 2 * auc**2 / (1 + auc)
        auc_variance = (
            auc * (1 - auc)
            + (positives - 1) * (q1 - auc**2)
            + (negatives - 1) * (q2 - auc**2)
        ) / (positives * negatives)
        auc_se = np.sqrt(max(0.0, auc_variance))
        auc_ci_low = max(0.0, auc - 1.96 * auc_se)
        auc_ci_high = min(1.0, auc + 1.96 * auc_se)
    grouped = pd.qcut(pd.Series(probabilities), q=min(10, len(probabilities)), duplicates="drop")
    metric_frame = pd.DataFrame({"response": response.to_numpy(), "probability": probabilities, "group": grouped})
    grouped_data = metric_frame.groupby("group", observed=False)
    observed = grouped_data["response"].sum()
    totals = grouped_data["response"].count()
    expected = grouped_data["probability"].sum()
    hl_statistic = float(
        (((observed - expected) ** 2) / expected.clip(lower=1e-12)
         + (((totals - observed) - (totals - expected)) ** 2)
         / (totals - expected).clip(lower=1e-12)).sum()
    ) if len(observed) > 1 else None
    hl_p_value = (
        float(chi2.sf(hl_statistic, len(observed) - 2))
        if hl_statistic is not None and len(observed) > 2
        else None
    )
    vifs: dict[str, Any] = {}
    from statsmodels.stats.outliers_influence import variance_inflation_factor

    vif_design = sm.add_constant(active, has_constant="add")
    for index, name in enumerate(vif_design.columns):
        if name != "const":
            try:
                vifs[name] = float(variance_inflation_factor(vif_design.to_numpy(), index))
            except (ValueError, np.linalg.LinAlgError):
                vifs[name] = float("inf")
    high_vifs = {name: value for name, value in vifs.items() if value >= 2.5}
    quality = {
        "Hosmer-Lemeshow statistic": hl_statistic,
        "Hosmer-Lemeshow p value": hl_p_value,
        "AUC-ROC": auc,
        "AUC-ROC 95% CI": f"{auc_ci_low:.2f}-{auc_ci_high:.2f}" if auc is not None else None,
        "AIC": float(fitted.aic),
        "VIF >= 2.5": "; ".join(f"{name}: {value:.2f}" for name, value in high_vifs.items()) or "None >= 2.5",
    }
    return results, omnibus, excluded, float(fitted.deviance), len(active.columns), quality


def format_p_value(value: float | None) -> Any:
    if value is None or not np.isfinite(value):
        return ""
    decimals = 2
    threshold = 0.01
    while value < threshold:
        decimals += 1
        threshold /= 10
    return f"{value:.{decimals}f}"


def format_result(result: tuple[float, float, float, float] | None) -> tuple[Any, Any, Any]:
    if result is None:
        return "", "", ""
    odds_ratio, lower, upper, p_value = result
    return f"{odds_ratio:.2f}", f"{lower:.2f}-{upper:.2f}", format_p_value(p_value)


def build_report(
    data: pd.DataFrame,
    response: str,
    predictors: list[str],
    references: dict[str, Any] | None = None,
    analysis_name: str = "All selected variables",
) -> pd.DataFrame:
    references = references or {}
    design, response_values, details = make_design(data, response, predictors, references)
    multivariable, _, multivariable_excluded, multivariable_deviance, multivariable_terms, quality = fit_logistic(
        design, response_values
    )
    rows: list[dict[str, Any]] = []
    quality_rows: list[dict[str, Any]] = []

    for predictor in predictors:
        predictor_data = data[[response, predictor]].dropna()
        predictor_design, predictor_response, predictor_details = make_design(
            predictor_data, response, [predictor], references
        )
        univariable, univariable_omnibus, univariable_excluded, _, _, univariable_quality = fit_logistic(
            predictor_design, predictor_response
        )
        if univariable_quality is not None and analysis_name == "All selected variables":
            quality_rows.append({"Model": f"Univariable: {predictor}", **univariable_quality})
        if predictor in predictor_design:
            reduced_design = design.drop(columns=[predictor], errors="ignore")
        else:
            reduced_design = design.drop(
                columns=[column for column in design if column.startswith(f"{predictor}=")],
                errors="ignore",
            )
        _, _, _, reduced_deviance, reduced_terms, _ = fit_logistic(reduced_design, response_values)
        multivariable_omnibus = None
        if multivariable_deviance is not None and reduced_deviance is not None:
            from scipy.stats import chi2

            degrees_of_freedom = multivariable_terms - reduced_terms
            if degrees_of_freedom > 0:
                multivariable_omnibus = float(
                    chi2.sf(reduced_deviance - multivariable_deviance, degrees_of_freedom)
                )
        info = details[predictor]
        levels = info["levels"]
        if info["numeric"]:
            level_rows = [("(continuous)", predictor)]
        else:
            level_rows = [(str(level), None if index == 0 else f"{predictor}={level}") for index, level in enumerate(levels)]

        for level, term in level_rows:
            level_mask = data[predictor].astype(str).eq(level) if level != "(continuous)" else data[predictor].notna()
            valid = data.loc[level_mask & data[response].notna(), response]
            n_positive = int((valid == 1).sum())
            denominator = int(valid.shape[0])
            univ_result = None if term is None or term in univariable_excluded else univariable.get(term)
            multi_result = None if term is None or term in multivariable_excluded else multivariable.get(term)
            univ_or, univ_ci, univ_p = format_result(univ_result)
            multi_or, multi_ci, multi_p = format_result(multi_result)
            if term is None and not info["numeric"]:
                univ_ci = multi_ci = "Reference"
            if term in univariable_excluded or term in multivariable_excluded:
                univ_ci = multi_ci = "Excluded (separation)"
            rows.append(
                {
                    "Explanatory variable": predictor,
                    "Level": level,
                    "n/N (%)": f"{n_positive}/{denominator} ({100 * n_positive / denominator:.1f}%)" if denominator else "0/0 (NA)",
                    "Univariable OR": univ_or,
                    "Univariable 95% CI": univ_ci,
                    "Univariable p value": univ_p,
                    "Univariable omnibus p value": format_p_value(univariable_omnibus) if level == levels[0] else "",
                    "Multivariable OR": multi_or,
                    "Multivariable 95% CI": multi_ci,
                    "Multivariable p value": multi_p,
                    "Multivariable omnibus p value": format_p_value(multivariable_omnibus) if level == levels[0] else "",
                }
            )
    report = pd.DataFrame(rows)
    report.insert(0, "Analysis", analysis_name)
    report.attrs["omnibus_p_values"] = {
        predictor: {
            "univariable": float(report.loc[
                report["Explanatory variable"].eq(predictor),
                "Univariable omnibus p value",
            ].iloc[0]) if report.loc[
                report["Explanatory variable"].eq(predictor),
                "Univariable omnibus p value",
            ].iloc[0] != "" else None,
            "multivariable": float(report.loc[
                report["Explanatory variable"].eq(predictor),
                "Multivariable omnibus p value",
            ].iloc[0]) if report.loc[
                report["Explanatory variable"].eq(predictor),
                "Multivariable omnibus p value",
            ].iloc[0] != "" else None,
        }
        for predictor in predictors
    }
    report.attrs["multivariable_p_values"] = {
        predictor: [
            multivariable[term][3]
            for term in (
                [predictor]
                if details[predictor]["numeric"]
                else [
                    f"{predictor}={level}"
                    for level in details[predictor]["levels"][1:]
                ]
            )
            if term in multivariable
        ]
        for predictor in predictors
    }
    report.attrs["quality_metrics"] = quality
    report.attrs["quality_rows"] = quality_rows + (
        [{"Model": analysis_name, **quality}] if quality is not None else []
    )
    return report


def qualifying_predictors(report: pd.DataFrame, predictors: list[str]) -> list[str]:
    omnibus = report.attrs.get("omnibus_p_values", {})
    return [
        predictor
        for predictor in predictors
        if any(
            p_value is not None and p_value < 0.20
            for p_value in omnibus.get(predictor, {}).values()
        )
    ]


def qualifying_predictors_strict(report: pd.DataFrame, predictors: list[str]) -> list[str]:
    """Keep variables with a significant omnibus or individual p-value."""
    omnibus = report.attrs.get("omnibus_p_values", {})
    individual = report.attrs.get("multivariable_p_values", {})
    qualifying: list[str] = []
    for predictor in predictors:
        omnibus_p = omnibus.get(predictor, {}).get("multivariable")
        individual_p_values = individual.get(predictor, [])
        if (
            omnibus_p is not None and omnibus_p < 0.05
        ) or any(p_value < 0.05 for p_value in individual_p_values):
            qualifying.append(predictor)
    return qualifying


def qualifying_predictors_omnibus_strict(
    report: pd.DataFrame, predictors: list[str]
) -> list[str]:
    """Keep variables whose multivariable omnibus p-value is below 0.05."""
    omnibus = report.attrs.get("omnibus_p_values", {})
    return [
        predictor
        for predictor in predictors
        if (
            omnibus.get(predictor, {}).get("multivariable") is not None
            and omnibus[predictor]["multivariable"] < 0.05
        )
    ]


def assemble_output(
    full_report: pd.DataFrame,
    model_reports: list[tuple[str, pd.DataFrame]],
) -> pd.DataFrame:
    descriptive_columns = ["Explanatory variable", "Level", "n/N (%)"]
    base = full_report[descriptive_columns].reset_index(drop=True)

    def model_block(
        source: pd.DataFrame,
        analysis_label: str,
        include_univariable: bool,
    ) -> pd.DataFrame:
        columns = ["Multivariable"]
        if include_univariable:
            columns = ["Univariable"]
        block_columns = [
            "Univariable OR" if include_univariable else "Multivariable OR",
            "Univariable 95% CI" if include_univariable else "Multivariable 95% CI",
            "Univariable p value" if include_univariable else "Multivariable p value",
            "Univariable omnibus p value"
            if include_univariable
            else "Multivariable omnibus p value",
        ]
        block = source[block_columns].copy()
        block.insert(0, "Analysis", analysis_label)
        block.columns = ["Analysis", "OR", "95% CI", "p value", "omnibus p value"]
        return block.reset_index(drop=True)

    blocks = [
        model_block(full_report, "Univariable", include_univariable=True),
        model_block(full_report, "Multivariable", include_univariable=False),
    ]
    keys = pd.MultiIndex.from_frame(base[["Explanatory variable", "Level"]])
    for analysis_label, source in model_reports:
        model = source.set_index(["Explanatory variable", "Level"])
        model = model.reindex(keys).reset_index(drop=True)
        blocks.append(
            model_block(
                model,
                analysis_label,
                include_univariable=False,
            )
        )
    return pd.concat([base, *blocks], axis=1)


def assemble_quality_report(reports: list[pd.DataFrame]) -> pd.DataFrame:
    columns = [
        "Model",
        "Hosmer-Lemeshow statistic",
        "Hosmer-Lemeshow p value",
        "AUC-ROC",
        "AUC-ROC 95% CI",
        "AIC",
        "Delta AIC",
        "VIF >= 2.5",
    ]
    rows = [row for report in reports for row in report.attrs.get("quality_rows", [])]
    quality = pd.DataFrame(rows)
    if quality.empty:
        return pd.DataFrame(columns=columns)
    quality["Delta AIC"] = quality["AIC"] - quality["AIC"].min()
    return quality[columns]


def main() -> None:
    import tkinter as tk
    from tkinter import messagebox

    path = choose_file()
    if path is None:
        return
    try:
        columns = [str(column).strip() for column in pd.read_excel(path, sheet_name="all", nrows=0).columns]
        response_selection = choose_columns(columns, "Select binary response variable", multiple=False)
        if not response_selection:
            return
        response = response_selection[0]
        predictor_columns = [column for column in columns if column != response]
        predictors = choose_columns(predictor_columns, "Select explanatory variables", multiple=True)
        if not predictors:
            return
        data = load_binary_data(path, response, predictors)
        references = choose_references(data, predictors)
        full_report = build_report(data, response, predictors, references)
        selected_predictors = qualifying_predictors(full_report, predictors)
        if selected_predictors:
            p20_report = build_report(
                data,
                response,
                selected_predictors,
                references,
                "Variables with omnibus p < 0.20",
            )
            model_reports: list[tuple[str, pd.DataFrame]] = [
                ("Multivariable (omnibus p < 0.20)", p20_report)
            ]
        else:
            model_reports = []
        current_report = p20_report if selected_predictors else None
        strict_predictors = (
            qualifying_predictors_strict(current_report, selected_predictors)
            if current_report is not None
            else []
        )
        strict_iteration = 1
        while strict_predictors and strict_predictors != selected_predictors:
            strict_label = (
                "Multivariable (omnibus or individual p < 0.05) "
                f"iteration {strict_iteration}"
            )
            strict_report = build_report(
                data,
                response,
                strict_predictors,
                references,
                strict_label,
            )
            model_reports.append(
                (
                    strict_label,
                    strict_report,
                )
            )
            current_report = strict_report
            remaining_predictors = qualifying_predictors_strict(
                strict_report, strict_predictors
            )
            if remaining_predictors == strict_predictors:
                break
            strict_predictors = remaining_predictors
            strict_iteration += 1

        omnibus_predictors = strict_predictors
        omnibus_iteration = 1
        while omnibus_predictors and current_report is not None:
            remaining_predictors = qualifying_predictors_omnibus_strict(
                current_report, omnibus_predictors
            )
            if remaining_predictors == omnibus_predictors:
                break
            omnibus_predictors = remaining_predictors
            if not omnibus_predictors:
                break
            omnibus_label = (
                "Multivariable (omnibus p < 0.05) "
                f"iteration {omnibus_iteration}"
            )
            omnibus_report = build_report(
                data,
                response,
                omnibus_predictors,
                references,
                omnibus_label,
            )
            model_reports.append(
                (
                    omnibus_label,
                    omnibus_report,
                )
            )
            current_report = omnibus_report
            omnibus_iteration += 1

        report = assemble_output(full_report, model_reports)
        quality_reports = [full_report, *[source for _, source in model_reports]]
        quality_report = assemble_quality_report(quality_reports)
        output_path = path.with_name(f"{path.stem}_logistic_regression.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            report.to_excel(writer, sheet_name="logistic_regression", index=False)
            quality_report.to_excel(writer, sheet_name="quality_metrics", index=False)
            worksheet = writer.sheets["logistic_regression"]
            from openpyxl.styles import Border, Side

            thick_line = Side(style="thick", color="000000")
            model_start_columns = {4, 9, *[14 + 5 * index for index in range(len(model_reports))]}
            for row_number in range(2, len(report) + 2):
                has_variable = bool(worksheet.cell(row_number, 1).value)
                is_last_level = (
                    has_variable
                    and (row_number == len(report) + 1
                    or worksheet.cell(row_number, 1).value
                    != worksheet.cell(row_number + 1, 1).value)
                )
                for column_number in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row_number, column_number)
                    cell.border = Border(
                        left=thick_line if column_number in model_start_columns else cell.border.left,
                        right=cell.border.right,
                        top=cell.border.top,
                        bottom=thick_line if is_last_level else cell.border.bottom,
                        diagonal=cell.border.diagonal,
                        diagonal_direction=cell.border.diagonal_direction,
                        diagonalUp=cell.border.diagonalUp,
                        diagonalDown=cell.border.diagonalDown,
                        outline=cell.border.outline,
                        vertical=cell.border.vertical,
                        horizontal=cell.border.horizontal,
                    )
            quality_worksheet = writer.sheets["quality_metrics"]
            for row_number in range(2, quality_worksheet.max_row + 1):
                for column_number in (2, 3, 4, 6, 7):
                    quality_worksheet.cell(row_number, column_number).number_format = "0.00"
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo("Regression complete", f"Results saved to:\n{output_path}")
        root.destroy()
    except Exception as error:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Regression failed", str(error))
        root.destroy()


if __name__ == "__main__":
    main()